from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_THEME_COLOR
from pathlib import Path

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'docs'/'presentation'/'LearnLens_AI_Judge_Deck.pptx'
OUT.parent.mkdir(parents=True,exist_ok=True)
prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
BLANK=prs.slide_layouts[6]
C={
 'dark':'103C3B','dark2':'183E3C','teal':'39756D','mint':'E5F1EA','mint2':'EDF5EF',
 'paper':'F7F5EF','white':'FFFFFF','muted':'758580','muted2':'61746E','gold':'A16B2B',
 'goldbg':'FFF5DF','rust':'A4473A','rustbg':'FFF0EC','line':'DBE4DC','pale':'F1F4EE','softdark':'B7D1C7'
}
def rgb(hexv): return RGBColor.from_string(hexv)
def bg(slide,color='paper'):
    fill=slide.background.fill; fill.solid(); fill.fore_color.rgb=rgb(C[color])
def rect(slide,x,y,w,h,fill, radius=False, line=None):
    sh=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb=rgb(C[fill])
    sh.line.color.rgb=rgb(C[line or fill])
    if radius:
        try: sh.adjustments[0]=0.12
        except Exception: pass
    return sh
def text(slide,x,y,w,h,txt,size=18,color='dark',bold=False,font='Aptos',align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, margin=0.04):
    sh=slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf=sh.text_frame; tf.clear(); tf.word_wrap=True
    tf.margin_left=Inches(margin); tf.margin_right=Inches(margin); tf.margin_top=Inches(margin); tf.margin_bottom=Inches(margin); tf.vertical_anchor=valign
    p=tf.paragraphs[0]; p.alignment=align
    run=p.add_run(); run.text=txt; run.font.name=font; run.font.size=Pt(size); run.font.bold=bold; run.font.color.rgb=rgb(C[color])
    return sh
def rich(slide,x,y,w,h,lines,fill=None,margin=0.14):
    if fill: rect(slide,x,y,w,h,fill,True)
    sh=slide.shapes.add_textbox(Inches(x+margin), Inches(y+margin), Inches(w-2*margin), Inches(h-2*margin))
    tf=sh.text_frame;tf.clear();tf.word_wrap=True;tf.margin_left=Inches(.02);tf.margin_right=Inches(.02);tf.margin_top=Inches(.02);tf.margin_bottom=Inches(.02)
    for i,item in enumerate(lines):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.text=item.get('text','');p.level=item.get('level',0);p.font.name='Aptos';p.font.size=Pt(item.get('size',15));p.font.bold=item.get('bold',False);p.font.color.rgb=rgb(C[item.get('color','dark')]);p.space_after=Pt(item.get('after',8));p.space_before=Pt(item.get('before',0))
    return sh
def title(slide,kicker,heading,sub=None):
    text(slide,.65,.35,12,.25,kicker.upper(),10,'gold',True)
    text(slide,.65,.7,12, .62, heading,29,'dark',True)
    if sub: text(slide,.68,1.38,11.9,.48,sub,13,'muted2')
def footer(slide,num):
    rect(slide,.65,7.13,12,.012,'line')
    text(slide,.68,7.18,6,.18,'LEARNLENS AI  |  HACKATHON DEMO',8,'muted',True)
    text(slide,12.05,7.18,.5,.18,f'{num:02d}',8,'muted',True,align=PP_ALIGN.RIGHT)
def pill(slide,x,y,w,label,fill='mint',color='teal'):
    rect(slide,x,y,w,.32,fill,True)
    text(slide,x+.07,y+.035,w-.14,.24,label,9,color,True,align=PP_ALIGN.CENTER,valign=MSO_ANCHOR.MIDDLE)
def arrow(slide,x1,y1,x2,y2,color='teal',width=1.5):
    s=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2));s.line.color.rgb=rgb(C[color]);s.line.width=Pt(width);s.line.end_arrowhead=True

# 1 Cover
s=prs.slides.add_slide(BLANK);bg(s,'dark')
rect(s,.72,.72,1.0,1.0,'teal',True)
text(s,.77,.93,.9,.5,'LL',22,'white',True,align=PP_ALIGN.CENTER,valign=MSO_ANCHOR.MIDDLE)
pill(s,2.0,.85,1.78,'AI FOR EDUCATION','mint','teal')
text(s,.75,1.75,7.5,1.5,'LearnLens AI',40,'white',True)
text(s,.78,3.1,7.5,1.0,'From wrong answers\nto a clearer next step.',25,'softdark',False)
text(s,.8,5.25,6.6,.7,'Diagnose  ·  Verify  ·  Teach  ·  Retest',15,'white',True)
text(s,.8,6.17,7.6,.5,'Craft N Code 2026  |  Judge presentation',11,'softdark')
# right-side loop illustration
rect(s,8.85,1.1,3.7,5.5,'dark2',True)
text(s,9.25,1.45,2.9,.45,'THE LEARNING LOOP',12,'softdark',True,align=PP_ALIGN.CENTER)
nodes=[('01','Work'),('02','Evidence'),('03','Gap'),('04','Verify'),('05','Lesson'),('06','Retest')]
for i,(num,label) in enumerate(nodes):
    x=9.25+(i%2)*1.55;y=2.1+(i//2)*1.25
    rect(s,x,y,1.35,.86,'teal' if i in [0,3,5] else 'dark',True)
    text(s,x+.08,y+.09,1.18,.2,num,9,'softdark',True,align=PP_ALIGN.CENTER)
    text(s,x+.08,y+.35,1.18,.28,label,12,'white',True,align=PP_ALIGN.CENTER)
text(s,8.95,6.25,3.5,.2,'Prototype - pilot evidence pending',9,'softdark',False,align=PP_ALIGN.CENTER)

# 2 Problem
s=prs.slides.add_slide(BLANK);bg(s);title(s,'01 / Problem','A score tells you what happened. Not always why.','LearnLens targets the moment after a student gets an answer wrong.')
rich(s,.7,2.15,3.8,3.45,[{'text':'FOR STUDENTS','size':11,'bold':True,'color':'gold','after':16},{'text':'“I lost a mark. What idea did I misunderstand?”','size':23,'bold':True,'after':16},{'text':'A correction may not show the underlying rule or the next practice step.','size':15,'color':'muted2'}],'white',.24)
rich(s,4.78,2.15,3.8,3.45,[{'text':'FOR TEACHERS','size':11,'bold':True,'color':'gold','after':16},{'text':'“Which misconception should I reteach first?”','size':23,'bold':True,'after':16},{'text':'Repeated class patterns need a quick, reviewable way to guide a short intervention.','size':15,'color':'muted2'}],'white',.24)
rich(s,8.86,2.15,3.75,3.45,[{'text':'THE PRODUCT IDEA','size':11,'bold':True,'color':'teal','after':16},{'text':'Turn work into a learning loop.','size':23,'bold':True,'after':16},{'text':'Evidence → a tentative gap → verification → a targeted lesson → a spaced retest.','size':15,'color':'muted2'}],'mint2',.24)
text(s,.75,6.0,11.8,.46,'Designed for students and teachers; not a replacement for teacher judgement or formal grading.',12,'muted2')
footer(s,2)

# 3 learning loop
s=prs.slides.add_slide(BLANK);bg(s);title(s,'02 / Workflow','A repeatable path from evidence to follow-up','The diagnosis is a hypothesis. A fresh question checks whether it is useful.')
steps=[('1','Upload work','PDF / image'),('2','Find patterns','Tags + evidence'),('3','Verify','Fresh questions'),('4','Teach','Retrieved notes'),('5','Practice','New variation'),('6','Retest','After 24 hours'),('7','Confirm','Provisional → confirmed')]
xs=[.72,2.52,4.32,6.12,7.92,9.72,11.52]
for i,(n,head,desc) in enumerate(steps):
    x=xs[i]
    rect(s,x,2.65,1.28,1.15,'dark' if i in [0,6] else 'white',True,'dark' if i in [0,6] else 'line')
    text(s,x+.08,2.77,1.1,.25,n,10,'gold' if i in [0,6] else 'teal',True,align=PP_ALIGN.CENTER)
    text(s,x+.05,3.08,1.18,.35,head,12,'white' if i in [0,6] else 'dark',True,align=PP_ALIGN.CENTER)
    text(s,x-.15,4.0,1.58,.5,desc,10,'muted2',False,align=PP_ALIGN.CENTER)
    if i < len(steps)-1: arrow(s,x+1.31,3.23,xs[i+1]-.04,3.23)
rect(s,1.05,5.1,11.2,.95,'goldbg',True)
text(s,1.3,5.34,10.7,.45,'One correct answer is not enough to confirm mastery. A successful retest must be due; early retests stay provisional.',14,'gold',True,align=PP_ALIGN.CENTER,valign=MSO_ANCHOR.MIDDLE)
footer(s,3)

# 4 architecture
s=prs.slides.add_slide(BLANK);bg(s);title(s,'03 / Engineering','A small architecture designed to degrade clearly','Core retrieval and deterministic fallback can run locally; external AI and persistence are optional.')
columns=[(.7,'STUDENT + TEACHER UI','Next.js / TypeScript',['Join by nickname + consent','Learning loop / read-aloud','Heatmap / CSV / reteach draft'],'mint2'),(4.05,'LEARNING API','FastAPI / Pydantic',['Work extraction + diagnosis','Misconception taxonomy','Spaced mastery + deletion'],'white'),(7.4,'CONTENT + RETRIEVAL','45 notes / 22 tags',['Local BM25-style ranking','Source IDs + snippets','No embedding API needed'],'white'),(10.0,'OPTIONAL SERVICES','Server-side only',['OpenRouter free route','Gemini optional fallback','Supabase + SQLite cache'],'goldbg')]
for x,head,sub,bullets,fill in columns:
    rect(s,x,2.2,2.7,3.35,fill,True)
    text(s,x+.18,2.45,2.3,.3,head,10,'gold' if fill!='mint2' else 'teal',True)
    text(s,x+.18,2.83,2.3,.33,sub,11,'dark',True)
    for j,b in enumerate(bullets):
        text(s,x+.19,3.4+j*.54,2.33,.44,'• '+b,11,'muted2')
for xx in [3.42,6.77,9.78]: arrow(s,xx,3.7,xx+.48,3.7)
text(s,.78,6.0,11.6,.48,'Secrets remain on the server. No live API or database credentials are included in this package.',12,'teal',True)
footer(s,4)

# 5 feature set
s=prs.slides.add_slide(BLANK);bg(s);title(s,'04 / What is built','Features that judges can click through','Implementation status is not the same as evidence of learning impact.')
features=[(.72,2.05,'01','Misconception diagnosis','22 taxonomy tags, rule fallback and a verification step.'), (6.85,2.05,'02','Grounded lessons','45 bundled notes; source IDs/snippets shown; English/Kannada/Hindi options.'),(.72,4.15,'03','Spaced mastery','BKT-inspired probability, 14-day half-life and delayed retest gate.'),(6.85,4.15,'04','Teacher decisions','Student × concept heatmap, top three tag patterns, 10-minute plan and CSV export.')]
for x,y,n,head,desc in features:
    rect(s,x,y,5.75,1.65,'white',True)
    rect(s,x+.2,y+.22,.58,.58,'mint',True)
    text(s,x+.22,y+.34,.54,.26,n,10,'teal',True,align=PP_ALIGN.CENTER)
    text(s,x+.95,y+.24,4.45,.34,head,17,'dark',True)
    text(s,x+.95,y+.7,4.45,.65,desc,12,'muted2')
text(s,.78,6.35,11.7,.35,'All diagnosis, translation and reteach suggestions need teacher review; mastery must not be used as an official grade.',11,'rust',True)
footer(s,5)

# 6 Retrieval/cache
s=prs.slides.add_slide(BLANK);bg(s);title(s,'05 / Trustworthy behaviour','Make the system explain where its help comes from','This is real local lexical retrieval over bundled notes, not a claim that a vector database is currently running.')
# pipeline cards
p=[(.78,'Student evidence','Question + answer'),(3.2,'Tag/rank','Rules + local BM25'),(5.62,'Retrieved notes','IDs + snippets'),(8.04,'AI (optional)','Provider A → B'),(10.46,'Learning response','Reviewable draft')]
for i,(x,h,sub) in enumerate(p):
    rect(s,x,2.5,2.0,1.25,'dark' if i in [0,4] else 'white',True,'dark' if i in [0,4] else 'line')
    text(s,x+.1,2.72,1.8,.3,h,12,'white' if i in [0,4] else 'dark',True,align=PP_ALIGN.CENTER)
    text(s,x+.1,3.1,1.8,.3,sub,10,'softdark' if i in [0,4] else 'muted2',False,align=PP_ALIGN.CENTER)
    if i<4: arrow(s,x+2.02,3.12,p[i+1][0]-.06,3.12)
rich(s,.8,4.42,3.72,1.33,[{'text':'Fallback, not fake certainty','size':14,'bold':True},{'text':'Unsupported offline topics return a clear availability warning.','size':11,'color':'muted2'}],'mint2',.17)
rich(s,4.8,4.42,3.72,1.33,[{'text':'Fingerprint cache','size':14,'bold':True},{'text':'Results are scoped to the student profile; deletion attempts to purge the profile cache.','size':11,'color':'muted2'}],'white',.17)
rich(s,8.8,4.42,3.72,1.33,[{'text':'Visible provenance','size':14,'bold':True},{'text':'Provider/cache status and retrieved note sources are inspectable in the app/API.','size':11,'color':'muted2'}],'white',.17)
text(s,.8,6.18,11.6,.3,'Library: 32 original notes + 13 AI-assisted drafts pending team/educator review; none are independently validated.',11,'gold',True)
footer(s,6)

# 7 safety/evidence
s=prs.slides.add_slide(BLANK);bg(s);title(s,'06 / Responsibility & evidence','We separate a working demo from a validated product','No surveys, pilot results, accuracy benchmark or before/after improvement numbers are invented.')
rect(s,.75,2.05,5.7,3.85,'mint2',True)
text(s,1.05,2.35,5.0,.38,'DESIGNED INTO THIS PROTOTYPE',13,'teal',True)
left=['Nickname-only class join; no student email/password','Consent before class membership is created','Privacy page: clear browser profile + request backend/cache deletion','Teacher key generated randomly; created key stored hashed','Provisional mastery and explicit human-review caveat','Synthetic heatmap and worksheet labeled as synthetic']
for i,line in enumerate(left):
    text(s,1.1,2.92+i*.45,4.98,.42,'✓  '+line,11,'dark')
rect(s,6.75,2.05,5.8,3.85,'goldbg',True)
text(s,7.05,2.35,5.15,.38,'PENDING BEFORE IMPACT CLAIMS',13,'gold',True)
right=['Student survey + teacher interviews','15–20 human-labelled worksheets; inter-rater agreement','Provider latency, cache hit and failure-rate logging','Paired pre/post test + delayed retest follow-up','Hindi/Kannada review with fluent speakers','Production security and deletion verification']
for i,line in enumerate(right):
    text(s,7.1,2.92+i*.45,5.0,.42,'○  '+line,11,'dark')
text(s,.82,6.25,11.8,.38,'A demo fixture is useful for QA; it is not a real student, a pilot, or a learning outcome.',12,'rust',True,align=PP_ALIGN.CENTER)
footer(s,7)

# 8 demo and close
s=prs.slides.add_slide(BLANK);bg(s,'dark')
text(s,.75,.55,11,.3,'07 / JUDGE WALKTHROUGH',10,'softdark',True)
text(s,.75,1.0,11.5,.78,'See the loop. Challenge the assumptions.',30,'white',True)
checklist=[('01','Join','Blue Falcon · CSE-A · consent'),('02','Analyse','Use the clearly synthetic PDF'),('03','Learn','Check note IDs, change language, read aloud'),('04','Retest','Show provisional state; explain 24-hour gate'),('05','Teach','Heatmap, top tags, plan draft and CSV')]
for i,(n,h,desc) in enumerate(checklist):
    x=.85+(i%3)*4.1; y=2.25+(i//3)*1.65
    rect(s,x,y,3.65,1.27,'dark2',True)
    text(s,x+.18,y+.16,.47,.25,n,10,'softdark',True)
    text(s,x+.72,y+.12,2.62,.32,h,16,'white',True)
    text(s,x+.72,y+.55,2.65,.52,desc,11,'softdark')
rect(s,.85,5.75,11.75,.76,'teal',True)
text(s,1.05,5.95,11.3,.34,'Next milestone: approved pilot + independently labelled worksheets + production privacy review.',14,'white',True,align=PP_ALIGN.CENTER)
text(s,.85,6.73,11.5,.22,'LearnLens AI  |  Prototype status and all evidence limitations are documented in README.md / docs/evidence.md',9,'softdark')
footer(s,8)

prs.core_properties.title='LearnLens AI - Hackathon Judge Deck'
prs.core_properties.subject='Adaptive learning prototype, implementation and honest evidence status'
prs.core_properties.author='LearnLens team'
prs.core_properties.keywords='education, misconception, retrieval, spaced retest, teacher dashboard'
prs.save(OUT)
print(OUT)
